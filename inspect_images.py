import docx
import sys

def map_images():
    doc2 = docx.Document('NHÓM 13_IS 425 E.docx')
    output = []
    output.append(f"Total paragraphs in Doc2: {len(doc2.paragraphs)}")
    
    img_count = 0
    for i, p in enumerate(doc2.paragraphs):
        xml = p._p.xml
        if 'graphic' in xml or 'blip' in xml:
            img_count += 1
            prev_txt = doc2.paragraphs[i-1].text.strip() if i > 0 else ''
            curr_txt = p.text.strip()
            next_txt = doc2.paragraphs[i+1].text.strip() if i < len(doc2.paragraphs)-1 else ''
            output.append(f"Img #{img_count} at P{i}: Text='{curr_txt}' | Prev='{prev_txt[:60]}' | Next='{next_txt[:60]}'")
            
    # Check tables for images
    tbl_img_count = 0
    for t_idx, t in enumerate(doc2.tables):
        for r_idx, r in enumerate(t.rows):
            for c_idx, c in enumerate(r.cells):
                if 'graphic' in c._tc.xml or 'blip' in c._tc.xml:
                    tbl_img_count += 1
                    cell_text = c.text.strip()
                    output.append(f"Tbl Img #{tbl_img_count} in Table {t_idx} Row {r_idx} Col {c_idx}: Text='{cell_text[:60]}'")
                    
    with open('doc2_images_map.txt', 'w', encoding='utf-8') as f:
        f.write('\n'.join(output))
    print(f"Total paragraph images: {img_count}, table images: {tbl_img_count}. Saved doc2_images_map.txt")

if __name__ == '__main__':
    map_images()
