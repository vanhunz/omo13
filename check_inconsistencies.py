# -*- coding: utf-8 -*-
import docx
import sys
import re

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('DoAnNhom13_IS425_OMO.docx')

print("=== CHECKING FOR INCONSISTENCIES IN PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt: continue
    
    # Check address
    if 'nguyễn thị thập' in txt.lower():
        print(f"P[{i}] WRONG ADDRESS: {txt}")
        
    # Check shipping / vouchers
    if any(k in txt.lower() for k in ['250k', '250.000', 'welcomeomo', 'omoship', 'omolove']):
        print(f"P[{i}] WRONG VOUCHER / SHIPPING: {txt}")

    # Check 4 categories vs 3 categories
    if '4 phân nhóm' in txt.lower() or 'phân nhóm 4' in txt.lower():
        print(f"P[{i}] WRONG CATEGORY COUNT: {txt}")

print("\n=== CHECKING FOR INCONSISTENCIES IN TABLES ===")
for t_idx, t in enumerate(doc.tables):
    for r_idx, r in enumerate(t.rows):
        for c_idx, c in enumerate(r.cells):
            txt = c.text.strip()
            if 'nguyễn thị thập' in txt.lower():
                print(f"Table {t_idx} Row {r_idx} Col {c_idx} WRONG ADDRESS: {txt}")
            if any(k in txt.lower() for k in ['250k', '250.000', 'welcomeomo', 'omoship', 'omolove']):
                print(f"Table {t_idx} Row {r_idx} Col {c_idx} WRONG VOUCHER / SHIPPING: {txt}")
            if '4 phân nhóm' in txt.lower():
                print(f"Table {t_idx} Row {r_idx} Col {c_idx} WRONG CATEGORY COUNT: {txt}")
