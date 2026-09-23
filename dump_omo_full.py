import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('DoAnNhom13_IS425_OMO.docx')

with open('omo_full_structure.txt', 'w', encoding='utf-8') as f:
    f.write(f"=== FULL DUMP OF DoAnNhom13_IS425_OMO.docx ===\n")
    f.write(f"Total Paragraphs: {len(doc.paragraphs)}\n")
    f.write(f"Total Tables: {len(doc.tables)}\n\n")
    
    # Iterate through body elements in order
    for elem in doc.element.body:
        tag = elem.tag.split('}')[-1]
        if tag == 'p':
            p = docx.text.paragraph.Paragraph(elem, doc)
            txt = p.text.strip()
            drawings = elem.xpath('.//w:drawing') + elem.xpath('.//w:pict')
            img_info = f" [IMGS: {len(drawings)}]" if drawings else ""
            if txt or drawings:
                f.write(f"P ({p.style.name}): {txt}{img_info}\n")
        elif tag == 'tbl':
            t = docx.table.Table(elem, doc)
            rows = len(t.rows)
            cols = len(t.columns) if rows > 0 else 0
            drawings = elem.xpath('.//w:drawing') + elem.xpath('.//w:pict')
            f.write(f"\n--- TABLE ({rows}x{cols}, IMGS: {len(drawings)}) ---\n")
            for r_idx, r in enumerate(t.rows):
                row_cells = []
                for c_idx, c in enumerate(r.cells):
                    c_drawings = len(c._tc.xpath('.//w:drawing') + c._tc.xpath('.//w:pict'))
                    c_img = f" [IMG:{c_drawings}]" if c_drawings else ""
                    row_cells.append(f"C{c_idx}: {c.text.strip().replace(chr(10), ' ')}{c_img}")
                f.write(f"  Row {r_idx}: {' | '.join(row_cells)}\n")
            f.write("----------------------------------------\n\n")

print("Dumped structure to omo_full_structure.txt")
