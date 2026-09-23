import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('DoAnNhom13_IS425_OMO.docx')

print(f"=== VERIFICATION OF DoAnNhom13_IS425_OMO.docx ===")
print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

# 1. Check for any leftover placeholders
print("\n--- 1. CHECKING FOR PLACEHOLDERS IN PARAGRAPHS ---")
placeholder_keywords = [
    "chèn ảnh", "chụp ảnh", "viết nội dung", "mô tả ở đây",
    "tham khảo:", "(ghi rõ", "...", "chưa có", "lưu ý là", "kẻ bảng"
]
found_placeholders = []
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    for kw in placeholder_keywords:
        if kw in t.lower() and not t.startswith("Hình ") and not t.startswith("Bảng "):
            found_placeholders.append((i, p.style.name, kw, t[:100]))

print(f"Found paragraph placeholders: {len(found_placeholders)}")
for idx, style, kw, txt in found_placeholders:
    print(f"  P[{idx}] ({style}) [kw: {kw}]: {txt}")

# 2. Check for empty/placeholder table cells
print("\n--- 2. CHECKING FOR EMPTY/PLACEHOLDER TABLE CELLS ---")
empty_cells_count = 0
for t_idx, t in enumerate(doc.tables):
    for r_idx, r in enumerate(t.rows):
        for c_idx, c in enumerate(r.cells):
            txt = c.text.strip()
            imgs = len(c._tc.xpath('.//w:drawing') + c._tc.xpath('.//w:pict'))
            if (txt == "" and imgs == 0) or any(kw in txt.lower() for kw in ["viết nội dung", "chèn ảnh", "mô tả ở đây"]):
                empty_cells_count += 1
                print(f"  Table {t_idx+1} [Row {r_idx}, Col {c_idx}]: \"{txt[:60]}\" (imgs: {imgs})")

print(f"Total problematic table cells: {empty_cells_count}")

# 3. Check Headings Structure
print("\n--- 3. CHECKING HEADINGS STRUCTURE ---")
headings = []
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if p.style.name.startswith('Heading') or any(t.startswith(prefix) for prefix in ['CHƯƠNG', '1.', '2.', '3.', '4.', '5.']):
        headings.append((i, p.style.name, t))

print(f"Total Headings: {len(headings)}")
for idx, style, txt in headings:
    print(f"  [{idx:3d}] ({style:15s}) {txt}")

# 4. Check Images
drawings_p = sum(len(p._p.xpath('.//w:drawing') + p._p.xpath('.//w:pict')) for p in doc.paragraphs)
drawings_t = sum(len(t._element.xpath('.//w:drawing') + t._element.xpath('.//w:pict')) for t in doc.tables)
print(f"\n--- 4. IMAGES COUNT ---")
print(f"Drawings in paragraphs: {drawings_p}")
print(f"Drawings in tables: {drawings_t}")
print(f"Total drawings: {drawings_p + drawings_t}")
