import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import re
import os

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell margins in dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Set subtle border for tables."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_cell_text(cell, text, bold=False, italic=False, size_pt=10.5, font_name="Times New Roman", color_rgb=(0,0,0), align=WD_ALIGN_PARAGRAPH.LEFT):
    """Format cell text cleanly."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(*color_rgb)

def insert_p_after(para, text, style='Normal', bold=False, italic=False, size_pt=12, color_rgb=(0,0,0), align=WD_ALIGN_PARAGRAPH.LEFT, space_before=3, space_after=3):
    """Insert a new formatted paragraph after a given paragraph."""
    new_p = OxmlElement('w:p')
    para._p.addnext(new_p)
    new_para = docx.text.paragraph.Paragraph(new_p, para._parent)
    try:
        new_para.style = style
    except:
        new_para.style = 'Normal'
    new_para.alignment = align
    new_para.paragraph_format.space_before = Pt(space_before)
    new_para.paragraph_format.space_after = Pt(space_after)
    new_para.paragraph_format.line_spacing = 1.25
    if text:
        run = new_para.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size_pt)
        run.font.color.rgb = RGBColor(*color_rgb)
    return new_para

def insert_table_after(para, rows, cols):
    """Create a new table and insert it right after the paragraph."""
    doc = para._parent
    new_tbl = doc.add_table(rows=rows, cols=cols)
    para._p.addnext(new_tbl._tbl)
    set_table_borders(new_tbl)
    new_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    return new_tbl

print("Loaded helper functions.")
