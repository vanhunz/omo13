import docx
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

def verify_report():
    filepath = "NHOM13_IS425_BAO_CAO_DO_AN_HOAN_CHINH.docx"
    print(f"=== VERIFYING FINAL REPORT: {filepath} ===")
    
    if not os.path.exists(filepath):
        print(f"ERROR: {filepath} does not exist!")
        return
        
    size_mb = os.path.getsize(filepath) / (1024 * 1024)
    print(f"File Size: {size_mb:.2f} MB")
    
    doc = docx.Document(filepath)
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")
    
    # 1. Check Chapters & Key Headings
    required_headings = [
        "CHƯƠNG I. TỔNG QUAN VỀ DỰ ÁN",
        "1.1. Giới thiệu tổng quan về thương hiệu",
        "1.1.1. Thành lập và ý tưởng dự án",
        "1.1.2. Lĩnh vực kinh doanh và danh mục sản phẩm",
        "1.2. Lý do lựa chọn ngành hàng",
        "1.3. Mô hình kinh doanh",
        "CHƯƠNG II. HOẠCH ĐỊNH CHIẾN LƯỢC MARKETING",
        "2.1. Phân khúc thị trường",
        "2.1.1. Khách hàng mục tiêu",
        "2.1.2. Đối thủ cạnh tranh",
        "2.2. Mô hình SWOT",
        "2.3. Mô hình 4P (Marketing Mix)",
        "2.4. Xác định KPIs",
        "CHƯƠNG III. THIẾT KẾ WEBSITE",
        "3.1. Phân tích hệ thống",
        "3.1.1 Yêu cầu hệ thống",
        "3.1.3. Demo website",
        "3.1.4. Tối ưu UI/UX",
        "3.2. Cài đặt Google Analytics",
        "3.3. Khai báo Google Search Console và Submit Sitemap",
        "3.4. Các Plugin đã cài đặt",
        "CHƯƠNG IV. SEO",
        "4.1. Nghiên cứu từ khóa",
        "4.2. Chiến lược nội dung",
        "4.2.1. Plan content",
        "4.2.2. Content tối ưu SEO",
        "4.3. SEO Onpage",
        "4.4. SEO Entity",
        "4.5. SEO Offpage",
        "CHƯƠNG V. ĐO LƯỜNG VÀ ĐÁNH GIÁ",
        "5.1. Audit Website",
        "5.1.1. SEOquake",
        "5.1.2. Screaming Frog",
        "5.1.3. Google Pagespeed Insights",
        "5.2. Đo lường hiệu quả",
        "5.2.1. Vị trí thứ hạng từ khóa",
        "5.2.3. Chỉ số Google Search Console",
        "5.2.4. Chỉ số Google Analytics",
        "5.3. Báo cáo KPIs",
        "5.4. Ưu điểm - Nhược điểm - Đề xuất giải pháp",
        "5.5. Bảng các công cụ sử dụng",
        "5.6. Bảng phân công nhiệm vụ và đánh giá thành viên"
    ]
    
    doc_text = "\n".join([p.text for p in doc.paragraphs])
    print("\n--- CHECKING REQUIRED HEADINGS ---")
    all_headings_found = True
    for h in required_headings:
        found = h.lower() in doc_text.lower()
        status = "✅ FOUND" if found else "❌ MISSING"
        print(f"[{status}] {h}")
        if not found:
            all_headings_found = False
            
    # 2. Check for leftover placeholders
    print("\n--- CHECKING LEFTOVER PLACEHOLDERS ---")
    suspicious_keywords = [
        "nhóm hạnh.pdf",
        "(yêu cầu ghi đủ",
        "(kẻ bảng 3 cột",
        "(chỉ ghi 10 link",
        "giảng viên yêu cầu bao nhiêu",
        "bài demo: link bài",
        "chụp ở seoquake"
    ]
    found_placeholders = []
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.lower()
        for sk in suspicious_keywords:
            if sk in txt:
                found_placeholders.append((f"P{i}", sk, p.text[:60]))
                
    for t_idx, t in enumerate(doc.tables):
        for r_idx, r in enumerate(t.rows):
            for c_idx, c in enumerate(r.cells):
                txt = c.text.lower()
                for sk in suspicious_keywords:
                    if sk in txt:
                        found_placeholders.append((f"Table {t_idx} R{r_idx}C{c_idx}", sk, c.text[:60]))
                        
    if found_placeholders:
        print(f"⚠️ Found {len(found_placeholders)} suspicious placeholder items:")
        for loc, sk, snippet in found_placeholders:
            print(f"  - {loc} ('{sk}'): {snippet}...")
    else:
        print("✅ ZERO placeholders found! All sections clean.")

    # 3. Check Images Count
    import zipfile
    with zipfile.ZipFile(filepath, 'r') as z:
        media_count = len([n for n in z.namelist() if n.startswith('word/media/')])
    print(f"\n--- IMAGES / MEDIA CHECK ---")
    print(f"✅ Total embedded images preserved: {media_count}")
    
    print("\n=== VERIFICATION RESULT: " + ("PERFECT / 100% PASS" if all_headings_found and not found_placeholders else "NEEDS REVIEW") + " ===")

if __name__ == '__main__':
    verify_report()
