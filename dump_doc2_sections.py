import docx

def dump_doc2():
    doc = docx.Document('NHÓM 13_IS 425 E.docx')
    with open('doc2_full_dump.txt', 'w', encoding='utf-8') as f:
        for i, p in enumerate(doc.paragraphs):
            img = '[IMG]' if ('graphic' in p._p.xml or 'blip' in p._p.xml) else ''
            f.write(f"P{i:03d} [{p.style.name}] {img}: {p.text}\n")
            
        f.write("\n" + "="*50 + "\nTABLES DUMP:\n")
        for t_idx, t in enumerate(doc.tables):
            f.write(f"\n--- Table {t_idx} ({len(t.rows)} rows x {len(t.columns)} cols) ---\n")
            for r_idx, r in enumerate(t.rows):
                row_str = " | ".join([c.text.strip().replace('\n', ' ') for c in r.cells])
                f.write(f"Row {r_idx}: {row_str}\n")

    print("Saved doc2_full_dump.txt")

if __name__ == '__main__':
    dump_doc2()
