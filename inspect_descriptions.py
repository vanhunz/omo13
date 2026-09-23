# -*- coding: utf-8 -*-
import docx
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('DoAnNhom13_IS425_OMO.docx')

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

# Dump paragraphs with index and text
with open('all_paragraphs_dump.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        f.write(f"--- P{i} [{p.style.name}] ---\n{p.text}\n\n")

# Dump all tables
with open('all_tables_dump.txt', 'w', encoding='utf-8') as f:
    for t_idx, t in enumerate(doc.tables):
        f.write(f"\n==================== TABLE {t_idx} (Rows: {len(t.rows)}, Cols: {len(t.columns)}) ====================\n")
        for r_idx, r in enumerate(t.rows):
            row_vals = [c.text.replace('\n', ' | ') for c in r.cells]
            f.write(f"Row {r_idx}: {row_vals}\n")

print("Dumped all_paragraphs_dump.txt and all_tables_dump.txt successfully!")
