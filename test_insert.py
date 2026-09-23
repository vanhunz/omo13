import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def insert_paragraph_after(paragraph, text=None, style=None):
    """Insert a new paragraph after the given paragraph."""
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    new_para = docx.text.paragraph.Paragraph(new_p, paragraph._parent)
    if text:
        new_para.text = text
    if style:
        new_para.style = style
    return new_para

def test_insertion():
    doc = docx.Document('NHÓM 13_IS 425 E.docx')
    print(f"Original paragraphs: {len(doc.paragraphs)}")
    print(f"Original tables: {len(doc.tables)}")

test_insertion()
