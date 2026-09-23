import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('DoAnNhom13_IS425_OMO.docx')

print(f"=== ALL TABLES IN DoAnNhom13_IS425_OMO.docx (Total: {len(doc.tables)}) ===")
for i, t in enumerate(doc.tables):
    rows = len(t.rows)
    cols = len(t.columns) if rows > 0 else 0
    drawings = len(t._element.xpath('.//w:drawing') + t._element.xpath('.//w:pict'))
    hdr = " | ".join([c.text.strip().replace('\n', ' ') for c in t.rows[0].cells]) if rows > 0 else ""
    print(f"\n--- Table {i+1} ({rows}x{cols}, imgs={drawings}) ---")
    print(f"Header: {hdr[:100]}")
    for r_idx in range(min(rows, 3)):
        r_txt = " | ".join([c.text.strip().replace('\n', ' ')[:40] for c in t.rows[r_idx].cells])
        print(f"  Row {r_idx}: {r_txt}")
