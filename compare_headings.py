import docx
import json
import re

def compare_file1_file2():
    doc1 = docx.Document('MauDoAnNhom_IS425.docx')
    doc2 = docx.Document('NHÓM 13_IS 425 E.docx')

    f1_headings = []
    for p in doc1.paragraphs:
        if p.style.name.startswith('Heading') or 'chương' in p.text.lower() or (p.text.strip() and p.text.strip()[0].isdigit() and '.' in p.text[:5]):
            f1_headings.append(p.text.strip())

    f2_headings = []
    for p in doc2.paragraphs:
        if p.style.name.startswith('Heading') or 'chương' in p.text.lower() or (p.text.strip() and p.text.strip()[0].isdigit() and '.' in p.text[:5]):
            f2_headings.append(p.text.strip())

    with open('headings_comparison.txt', 'w', encoding='utf-8') as f:
        f.write("=== FILE 1 HEADINGS (TEMPLATE / REQUIREMENTS) ===\n")
        for h in f1_headings:
            f.write(h + "\n")
        f.write("\n" + "="*50 + "\n")
        f.write("=== FILE 2 HEADINGS (CURRENT DRAFT) ===\n")
        for h in f2_headings:
            f.write(h + "\n")

    print("Saved headings_comparison.txt")

if __name__ == '__main__':
    compare_file1_file2()
