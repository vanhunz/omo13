import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('DoAnNhom13_IS425_OMO.docx')

print("--- ALL IMAGES IN PARAGRAPHS ---")
for i, p in enumerate(doc.paragraphs):
    drawings = p._p.xpath('.//w:drawing') + p._p.xpath('.//w:pict')
    if drawings:
        prev_txt = doc.paragraphs[i-1].text.strip()[:60] if i > 0 else ""
        next_txt = doc.paragraphs[i+1].text.strip()[:60] if i < len(doc.paragraphs)-1 else ""
        print(f"P[{i}] has {len(drawings)} imgs | Prev: \"{prev_txt}\" | Next: \"{next_txt}\"")

print("\n--- ALL IMAGES IN TABLES ---")
for t_idx, t in enumerate(doc.tables):
    for r_idx, r in enumerate(t.rows):
        for c_idx, c in enumerate(r.cells):
            drawings = c._tc.xpath('.//w:drawing') + c._tc.xpath('.//w:pict')
            if drawings:
                c_txt = c.text.strip()[:40]
                print(f"Table {t_idx+1} [R{r_idx} C{c_idx}] has {len(drawings)} imgs | Text: \"{c_txt}\"")
